from django.shortcuts import render, redirect
from django.contrib import messages
from .models import *
from .forms import *
from t_groupe.models import Groupe
from django.contrib.auth.decorators import login_required
from django.db import transaction, IntegrityError
from django.db.models.deletion import ProtectedError
from django_tenants.utils import get_tenant_model, schema_context
from django.utils import timezone
from .sync_utils import *
from django.http import JsonResponse, HttpResponse
import csv
import openpyxl
from django.core.exceptions import ValidationError
from django.db.models import Q, Prefetch
from .import_utils import handle_uploaded_file, verify_data, import_data
import json
from django.core.paginator import Paginator
from django.db.models import Q

import json
from .data.modules_data import MODULES_DATA

from app.models import Institut
from t_crm.models import UserActionLog

def listModules(request):
    modules = Modules.objects.all()
    return render(request, 't_formations/modules.html', {'modules': modules})

def listSpecialites(request):
    queryset = Specialites.objects.all()
    if request.tenant.tenant_type != 'master':
        queryset = queryset.filter(is_visible=True)
    
    specialites = queryset.order_by('formation')
    context = {
        'liste' : specialites,
        'tenant' : request.tenant
    }
    return render(request, 'tenant_folder/formations/liste_des_specialites.html', context)

@login_required(login_url="institut_app:login")
def listFormations(request):
    # Optimisation avec prefetch_related pour éviter les requêtes N+1
    queryset = Specialites.objects.all()
    if request.tenant.tenant_type != 'master':
        queryset = queryset.filter(is_visible=True)

    formations = Formation.objects.all().prefetch_related(
        Prefetch('formation_specilite', queryset=queryset, to_attr='visible_specialites'),
        'visible_specialites__modules_set',
        'dossierinscription_set'
    )
    
    any_missing = False
    for f in formations:
        f.missing_info = []
        # Vérification des champs de base
        if not f.type_formation: f.missing_info.append("Type")
        if not f.prix_formation or f.prix_formation == 0: f.missing_info.append("Prix")
        if not f.frais_inscription or f.frais_inscription == 0: f.missing_info.append("Frais")
        if not f.duree: f.missing_info.append("Durée")
        if not f.partenaire: f.missing_info.append("Partenaire")
        if not f.entite_legal: f.missing_info.append("Entité")
        if not f.description: f.missing_info.append("Description")
        if not f.qualification: f.missing_info.append("Qualification")
        
        # Vérification du Dossier d'Inscription
        if not f.dossierinscription_set.exists():
            f.missing_info.append("Dossier")

        # Vérification de l'architecture (Spécialités et Modules)
        specs = f.visible_specialites
        if not specs:
            f.missing_info.append("Spécialités")
        else:
            for s in specs:
                if not s.modules_set.exists():
                    f.missing_info.append(f"Modules ({s.label})")
        
        if f.missing_info:
            any_missing = True

    context = {
        'liste': formations,
        'tenant' : request.tenant,
        'has_any_missing': any_missing
    }
    return render(request, 'tenant_folder/formations/liste_des_formations.html', context)

def listFraisInscription(request):
    frais = FraisInscription.objects.all()
    return render(request, 't_formations/frais.html', {'frais': frais})

def ListeDesInstituts(request):
    liste = Institut.objects.filter(is_visible=True).exclude(Q(schema_name='public') | Q(tenant_type='master'))
    context = {
        'liste' : liste,
        'tenant' : request.tenant,
    }
    return render(request, 'tenant_folder/formations/liste_des_instituts.html', context)

@login_required(login_url="institut_app:login")
@transaction.atomic
def addFormation(request):
    form = NewFormationFormMaster(current_tenant = request.tenant)
    if request.method == 'POST':
        form = NewFormationFormMaster(request.POST, current_tenant =request.tenant)
        if form.is_valid():
            form.save()
            messages.success(request, 'Formation ajoutée avec succès')
            return redirect('t_formations:listFormations')
        else:
            messages.error(request, 'Une erreur s\'est produite lors du traitement de la requête')
            return redirect('t_formations:addFormation')

    context = {
        'form' : form,
        'tenant' : request.tenant,
    }
    return render(request, 'tenant_folder/formations/nouvelle_formations.html', context)

def ApiCheckIfFormationCompleted(request):
    code = request.GET.get('code_formation')
    if not code:
        return JsonResponse({'error': 'Missing code_formation'}, status=400)

    formation = Formation.objects.filter(code=code).first()
    if not formation:
        return JsonResponse({'error': 'Formation not found'}, status=404)

    specialites = Specialites.objects.filter(formation=formation)

    if not specialites.exists():
        return JsonResponse({'completed': False, 'reason': 'Aucune spécialité trouvée <br> Veuillez ajouter au moins une spécialité.'})

    # Vérifier que chaque spécialité a au moins un module
    for specialite in specialites:
        if not Modules.objects.filter(specialite=specialite).exists():
            return JsonResponse({'completed': False, 'reason': f'<h5>la spécialité {specialite.label} n\'a aucun module </h5><br><br> <p> Veuillez ajouter au moins un module avant de pouvoir synchroniser la formation.</p>'})

    # Si tout est bon
    return JsonResponse({'completed': True})


@login_required(login_url='institut_app:login')
@transaction.atomic
def updateFormation(request, pk):
    
    formation = Formation.objects.get(pk = pk)
    form = NewFormationFormMaster(instance = formation)
    if request.method == "POST":
        form = NewFormationFormMaster(request.POST, instance=formation)
        if form.is_valid():
            updated_formation = form.save()
            updated_formation.updated = True
            updated_formation.save()

            messages.success(request, 'Les informations de la formation ont été modifiées avec succès')
            return redirect('t_formations:listFormations')
        else:
            messages.error(request, 'Une erreur s\'est produite lors du traitement de la requête')
            return redirect('t_formations:updateFormation', pk)
    else:
        context = {
            'form' : form,
            'tenant' : request.tenant
        }
        return render(request, 'tenant_folder/formations/update_formation.html', context)

@transaction.atomic
@login_required(login_url='institut_app:login')
def AddPartenaire(request):
    if request.method == 'POST':
        form = NewPartenaireForm(request.POST, request.FILES, current_tenant=request.tenant)
        if form.is_valid():
            partenaire = form.save()
            UserActionLog.objects.create(
                user=request.user,
                action_type='CREATE',
                target_model='Partenaire',
                target_id=str(partenaire.id),
                details=f"Création du partenaire : {partenaire.nom} (Code: {partenaire.code})",
                ip_address=request.META.get('REMOTE_ADDR')
            )
            messages.success(request, 'Partenaire ajouté avec succès')
            return redirect('t_formations:listPartenaires')
    else:
        form = NewPartenaireForm(current_tenant=request.tenant)

    context = {
        'form': form,
        'tenant': request.tenant,
    }
    return render(request, 'tenant_folder/formations/new_partenaire.html', context)

def detailsPartenaire(request, pk):
    partenaire = Partenaires.objects.get(id=pk)

    context = {
        'partenaire': partenaire,           
        'tenant': request.tenant,
    }
    return render(request, 'tenant_folder/formations/details_partenaire.html', context)

def ApiGetPartenaireSync(request):
    code = request.GET.get('code_partenaire')
    liste = Institut.objects.filter(is_visible=True).exclude(Q(schema_name='public') | Q(tenant_type='master'))
    instituts = []

    for institut in liste:
        with schema_context(institut.schema_name):
            has_partenaire = Partenaires.objects.filter(code=code).exists()

        instituts.append({
            'schema_name': institut.schema_name,
            'has_partenaire': has_partenaire,
        })

    return JsonResponse(instituts, safe=False)

def ApigetFormationSync(request):
    code = request.GET.get('code_formation')
    liste = Institut.objects.filter(is_visible=True).exclude(Q(schema_name='public') | Q(tenant_type='master'))
    formation = Formation.objects.get(code=code)

    instituts = []

    for institut in liste:
        with schema_context(institut.schema_name):
            has_formation = Formation.objects.filter(code=code).exists()
            has_partenaire = False
            if formation.partenaire:
                has_partenaire = Partenaires.objects.filter(code=formation.partenaire.code).exists()

        instituts.append({
            'schema_name': institut.schema_name,
            'nom': institut.nom,
            'has_formation': has_formation,
            'has_partenaire': has_partenaire,
        })

    return JsonResponse(instituts, safe=False)

@login_required(login_url="institut_app:login")
def ApiUnsyncFormation(request):
    if request.method != "POST":
        return JsonResponse({'status': False, 'message': 'Méthode non autorisée'})
    
    code_formation = request.POST.get('code_formation')
    schema_name = request.POST.get('schema_name')
    
    if not code_formation or not schema_name:
        return JsonResponse({'status': False, 'message': 'Paramètres manquants'})

    try:
        institut = Institut.objects.get(schema_name=schema_name)
        with schema_context(institut.schema_name):
            try:
                formation_local = Formation.objects.get(code=code_formation)
                
                # 1. Delete DossierInscription (has on_delete=models.DO_NOTHING in models)
                DossierInscription.objects.filter(formation=formation_local).delete()
                
                # 2. Delete Specialities (and their modules via CASCADE)
                specialites = Specialites.objects.filter(formation=formation_local)
                
                # Delete DoubleDiplomation that reference these specialities
                DoubleDiplomation.objects.filter(Q(specialite1__in=specialites) | Q(specialite2__in=specialites)).delete()
                
                specialites.delete()
                
                # 3. Delete the formation itself
                formation_local.delete()
                
                return JsonResponse({'status': True, 'message': f'La formation a été retirée de {institut.nom}.'})
            except Formation.DoesNotExist:
                return JsonResponse({'status': False, 'message': 'Formation introuvable dans ce tenant.'})
    except Institut.DoesNotExist:
        return JsonResponse({'status': False, 'message': 'Institut introuvable.'})
    except ProtectedError:
        return JsonResponse({'status': False, 'message': 'Suppression impossible : des données actives (promos, inscriptions) sont liées à cette formation dans ce tenant.'})
    except Exception as e:
        return JsonResponse({'status': False, 'message': str(e)})

@login_required(login_url="institut_app:login")
def ApiClearTenantModules(request):
    if request.tenant.tenant_type != 'master':
        return JsonResponse({'status': False, 'message': 'Action non autorisée'})

    if request.method != "POST":
        return JsonResponse({'status': False, 'message': 'Méthode non autorisée'})

    specialite_code = request.POST.get('specialite_code')
    schema_name = request.POST.get('schema_name')

    if not specialite_code or not schema_name:
        return JsonResponse({'status': False, 'message': 'Paramètres manquants'})

    try:
        institut = Institut.objects.get(schema_name=schema_name)
        with schema_context(institut.schema_name):
            try:
                # Find the speciality in the target schema by its code
                specialite_local = Specialites.objects.get(code=specialite_code)
                # Delete all modules associated with this speciality in the target schema
                Modules.objects.filter(specialite=specialite_local).delete()
                return JsonResponse({'status': True, 'message': f'La liste des modules a été vidée pour {institut.nom}.'})
            except Specialites.DoesNotExist:
                return JsonResponse({'status': False, 'message': 'Spécialité introuvable dans ce tenant.'})
    except Institut.DoesNotExist:
        return JsonResponse({'status': False, 'message': 'Institut introuvable.'})
@login_required(login_url="institut_app:login")
def ApiSyncSpecialiteModulesToTenant(request):
    if request.tenant.tenant_type != 'master':
        return JsonResponse({'status': False, 'message': 'Action non autorisée'})

    if request.method != "POST":
        return JsonResponse({'status': False, 'message': 'Méthode non autorisée'})

    spec_code = request.POST.get('spec_code')
    schema_name = request.POST.get('schema_name')

    if not spec_code or not schema_name:
        return JsonResponse({'status': False, 'message': 'Paramètres manquants'})

    try:
        specialite_master = Specialites.objects.get(code=spec_code)
        institut = Institut.objects.get(schema_name=schema_name)
        
        # Ensure formation and speciality are up to date in the target tenant
        with schema_context(institut.schema_name):
            try:
                sync_formation = Formation.objects.get(code=specialite_master.formation.code)
            except Formation.DoesNotExist:
                if specialite_master.formation.type_formation == 'etrangere':
                    sync_formation = update_or_create_formation_in_tenant(specialite_master.formation, institut.schema_name)
                else:
                    return JsonResponse({'status': False, 'message': 'La formation doit d\'abord être synchronisée dans cet établissement.'})

        # Update speciality info (label, price, etc.)
        sync_specialite = update_or_create_specialite_in_tenant(specialite_master, sync_formation, institut.schema_name)

        # Get modules from master
        modules_master = Modules.objects.filter(specialite=specialite_master)
        
        # Sync each module
        for module in modules_master:
            update_or_create_module_in_tenant(module, sync_specialite, institut.schema_name)
            
        return JsonResponse({'status': True, 'message': f'La spécialité et ses modules ont été synchronisés avec succès pour {institut.nom}.'})
    except Specialites.DoesNotExist:
        return JsonResponse({'status': False, 'message': 'Spécialité source introuvable.'})
@login_required(login_url="institut_app:login")
def ApiCheckSpecialiteSyncState(request):
    if request.tenant.tenant_type != 'master':
        return JsonResponse({'status': False, 'message': 'Action non autorisée'})

    spec_code = request.GET.get('spec_code')
    if not spec_code:
        return JsonResponse({'status': False, 'message': 'Paramètre manquant'})

    try:
        specialite_master = Specialites.objects.get(code=spec_code)
        formation_master = specialite_master.formation
        modules_master = Modules.objects.filter(specialite=specialite_master)
        
        if formation_master.type_formation == 'etrangere':
            liste = Institut.objects.filter(is_visible=True, tenant_type='second').exclude(schema_name='public')
        else:
            liste = Institut.objects.filter(is_visible=True).exclude(Q(schema_name='public') | Q(tenant_type='master'))

        results = []
        for institut in liste:
            state = {
                'nom': institut.nom,
                'schema_name': institut.schema_name,
                'spec_exists': False,
                'spec_up_to_date': True,
                'modules_master': modules_master.count(),
                'modules_tenant': 0,
                'diff_found': False,
                'message': ''
            }
            
            with schema_context(institut.schema_name):
                try:
                    spec_tenant = Specialites.objects.get(code=spec_code)
                    state['spec_exists'] = True
                    state['is_visible_tenant'] = spec_tenant.is_visible
                    
                    # Compare ALL fields
                    diffs = []
                    if spec_tenant.label != specialite_master.label: diffs.append("libellé")
                    if spec_tenant.duree != specialite_master.duree: diffs.append("durée")
                    if spec_tenant.version != specialite_master.version: diffs.append("version")
                    if spec_tenant.condition_access != specialite_master.condition_access: diffs.append("conditions")
                    if spec_tenant.nb_semestre != specialite_master.nb_semestre: diffs.append("nb sem.")
                    if spec_tenant.branche != specialite_master.branche: diffs.append("branche")
                    if spec_tenant.abr != specialite_master.abr: diffs.append("abr.")
                    if spec_tenant.nb_tranche != specialite_master.nb_tranche: diffs.append("tranches")
                    if not specialite_master.is_visible and spec_tenant.is_visible:
                        diffs.append("visibilité (désactivé)")
                    
                    if diffs:
                        state['spec_up_to_date'] = False
                        state['diff_found'] = True
                        state['message'] = f"Différences: {', '.join(diffs)}"
                    
                    modules_tenant = Modules.objects.filter(specialite=spec_tenant)
                    state['modules_tenant'] = modules_tenant.count()
                    
                    if state['modules_tenant'] != state['modules_master']:
                        state['diff_found'] = True
                        if not state['message']:
                            state['message'] = "Nombre de modules différent"
                        else:
                            state['message'] += " + Nb modules"
                            
                except Specialites.DoesNotExist:
                    state['spec_exists'] = False
                    state['diff_found'] = True
                    state['message'] = "Spécialité manquante"
            
            results.append(state)
            
        return JsonResponse({'status': True, 'results': results})
    except Exception as e:
        return JsonResponse({'status': False, 'message': str(e)})

@login_required(login_url="institut_app:login")
def ApiSyncUpdateSpecialite(request):
    if request.tenant.tenant_type != 'master':
        return JsonResponse({'status': False, 'message': 'Action non autorisée'})

    if request.method != "POST":
        return JsonResponse({'status': False, 'message': 'Méthode non autorisée'})

    spec_code = request.POST.get('spec_code')
    
    if not spec_code:
        return JsonResponse({'status': False, 'message': 'Paramètre manquant (spec_code)'})

    try:
        specialite_master = Specialites.objects.get(code=spec_code)
        formation_master = specialite_master.formation
        
        # Determine target tenants
        if formation_master.type_formation == 'etrangere':
            liste = Institut.objects.filter(is_visible=True, tenant_type='second').exclude(schema_name='public')
        else:
            liste = Institut.objects.filter(is_visible=True).exclude(Q(schema_name='public') | Q(tenant_type='master'))

        updated_count = 0
        errors = []
        
        for institut in liste:
            try:
                with schema_context(institut.schema_name):
                    # For foreign, we create/update. For national, only if exists.
                    if formation_master.type_formation != 'etrangere' and not Specialites.objects.filter(code=spec_code).exists():
                        continue
                    
                    # Ensure formation exists in tenant
                    try:
                        sync_formation = Formation.objects.get(code=formation_master.code)
                    except Formation.DoesNotExist:
                        if formation_master.type_formation == 'etrangere':
                            sync_formation = update_or_create_formation_in_tenant(formation_master, institut.schema_name)
                        else:
                            continue

                    # Sync Speciality
                    sync_specialite = update_or_create_specialite_in_tenant(specialite_master, sync_formation, institut.schema_name)
                    
                    # Sync Modules
                    modules_master = Modules.objects.filter(specialite=specialite_master)
                    for module in modules_master:
                        update_or_create_module_in_tenant(module, sync_specialite, institut.schema_name)
                    
                    updated_count += 1
            except Exception as e:
                errors.append(f"{institut.nom}: {str(e)}")

        message = f"La spécialité a été mise à jour dans {updated_count} instituts."
        if errors:
            message += " Quelques erreurs : " + ", ".join(errors[:3])

        return JsonResponse({'status': True, 'message': message})
        
    except Specialites.DoesNotExist:
        return JsonResponse({'status': False, 'message': 'Spécialité source introuvable.'})
    except Exception as e:
        return JsonResponse({'status': False, 'message': str(e)})

@login_required(login_url="institut_app:login")
def ApiLoadFormations(request):
    if request.method == "GET":
        liste = Formation.objects.all().values('id','nom')
        return JsonResponse(list(liste), safe=False)
    else:
        return JsonResponse({'status':"error"})
    

@login_required(login_url="institut_app:login")
def ApiLoadSpecialites(request):
    if request.method == "GET":
        queryset = Specialites.objects.all()
        if request.tenant.tenant_type != 'master':
            queryset = queryset.filter(is_visible=True)
            
        liste = queryset.values('id','code','label','version','formation__id')
        return JsonResponse(list(liste), safe=False)
    else:
        return JsonResponse({'status':"error"})



@login_required(login_url="institut_app:login")
@transaction.atomic
def ApiImportModulesSpecialite(request):
    if request.method != "POST":
        return JsonResponse({'status': 'error', 'message': 'Méthode non autorisée'})

    id_specialite = request.POST.get('id_specialite')
    
    try:
        specialite = Specialites.objects.get(id=id_specialite)
    except Specialites.DoesNotExist:
        return JsonResponse({'status': 'error', 'message': 'Spécialité introuvable'})

    # Find matching data
    spec_data = None
    for item in MODULES_DATA:
        # Match by code or label (case insensitive)
        if (item.get('code') and item.get('code') == specialite.code) or \
           (item.get('specialite').lower() == specialite.label.lower()):
            spec_data = item
            break
    
    if not spec_data:
        return JsonResponse({'status': 'error', 'message': f'Aucun module par défaut trouvé pour cette spécialité: {specialite.label}'})

    count_created = 0
    count_updated = 0

    for module_data in spec_data.get('modules', []):
        code_interne = module_data.get('code_interne') or module_data.get('code')
        designation = module_data.get('designation')
        coef = module_data.get('coef')
        duree= module_data.get('duree')

        if not code_interne:
            continue

        module, created = Modules.objects.update_or_create(
            specialite=specialite,
            code_interne=code_interne,
            defaults={
                'label': designation,
                'coef' : coef,
                'duree' : duree,
                'created_by': request.user
            }
        )
        
        # Ensure code is generated if empty
        if created and not module.code:
             module.save() # Triggers generate_code in save()

        if created:
            count_created += 1
        else:
            count_updated += 1

    return JsonResponse({
        'status': 'success', 
        'message': f'Importation terminée: {count_created} créés, {count_updated} mis à jour.',
        'created': count_created
    })

######### methode de syncronisation des formations #############################################

##### Synchronisation des formations et spécialités dans un tenant spécifique ##################
def ApiSyncFormation(request):
    code_formation = request.POST.get('code_formation')
    schema_name = request.POST.get('schema_name')

    try:
        formation = Formation.objects.get(code=code_formation)
        institut = Institut.objects.get(schema_name=schema_name)
        master_tenant = Institut.objects.filter(tenant_type='master').exclude(schema_name='public').first()

        sync_formation_to_tenant(formation, institut, master_tenant)

        return JsonResponse({'status': True, 'message': 'Formation, dossier, spécialités et modules synchronisés avec succès'})
    except Formation.DoesNotExist:
        return JsonResponse({'status': False, 'message': 'Formation introuvable'})
    except Institut.DoesNotExist:
        return JsonResponse({'status': False, 'message': 'Institut introuvable'})
    except Exception as e:
        return JsonResponse({'status': False, 'message': str(e)})
##### Synchronisation des formations et spécialités dans un tenant spécifique ##################

@login_required(login_url="institut_app:login")
def ApiSyncUpdateFormation(request):
    code_formation = request.POST.get('code_formation')
    
    try:
        formation = Formation.objects.get(code=code_formation)
        master_tenant = Institut.objects.filter(tenant_type='master').exclude(schema_name='public').first()
    except Formation.DoesNotExist:
        return JsonResponse({'status': False, 'message': 'Formation introuvable.'})

    # Filter: Foreign formations are only for 'second' tenants
    # National formations can be for both 'second' and 'associe'
    if formation.type_formation == 'etrangere':
        liste = Institut.objects.filter(is_visible=True, tenant_type='second').exclude(schema_name='public')
    else:
        liste = Institut.objects.filter(is_visible=True).exclude(Q(schema_name='public') | Q(tenant_type='master'))

    formation.updated = False
    formation.save()
    
    updated_count = 0
    try:
        for institut in liste:
            with schema_context(institut.schema_name):
                # For foreign formations, we always update or create.
                # For others, we only update if it already exists in the tenant.
                if formation.type_formation != 'etrangere' and not Formation.objects.filter(code=formation.code).exists():
                    continue
                    
            sync_formation_to_tenant(formation, institut, master_tenant)
            updated_count += 1

        return JsonResponse({
            'status': True, 
            'message': f'La formation a été mise à jour dans {updated_count} instituts.'
        })
    except Exception as e:
        return JsonResponse({'status': False, 'message': f"Erreur lors de la mise à jour globale: {str(e)}"})

##### Synchronisation (Modification et création) des formations et spécialités dans tous les tenants ##################

def ApiCheckFormationState(request):
    code_formation = request.GET.get('code_formation')
    formation = Formation.objects.get(code=code_formation)

    if formation.updated:
        return JsonResponse({'status': True})
    else:
        return JsonResponse({'status': False}) 

def ApiSyncPartenaire(request):
    code_partenaire = request.POST.get('code_partenaire')
    shema_name = request.POST.get('schema_name')
    partenaire = Partenaires.objects.get(code=code_partenaire)

    institut = Institut.objects.get(schema_name=shema_name)
    with schema_context(institut.schema_name):
        sync_partenaire, created = Partenaires.objects.update_or_create(
            code=partenaire.code,
            defaults={
                'nom': partenaire.nom,
                'adresse': partenaire.adresse,
                'telephone': partenaire.telephone,
                'email': partenaire.email,
                'site_web': partenaire.site_web,
                'type_partenaire': partenaire.type_partenaire,
            }
        )
        UserActionLog.objects.create(
            user=request.user,
            action_type='UPDATE',
            target_model='Partenaire',
            target_id=str(partenaire.id),
            details=f"Synchronisation du partenaire {partenaire.nom} vers le tenant {institut.nom}",
            ip_address=request.META.get('REMOTE_ADDR')
        )
        message = 'Partenaire synchronisé avec succès' if created else 'Informations du partenaire mises à jour avec succès'
        return JsonResponse({'success': True, 'message': message})
        
def deletePartenaire(request, pk):
    partenaire = Partenaires.objects.get(id=pk)
    details = f"Suppression du partenaire : {partenaire.nom} (Code: {partenaire.code})"
    partenaire_id = partenaire.id
    partenaire.delete()
    UserActionLog.objects.create(
        user=request.user,
        action_type='DELETE',
        target_model='Partenaire',
        target_id=str(partenaire_id),
        details=details,
        ip_address=request.META.get('REMOTE_ADDR')
    )
    messages.success(request, 'Partenaire supprimé avec succès')
    return redirect('t_formations:listPartenaires')

@login_required(login_url="institut_app:login")
def ListeDesPartenaires(request):
    liste = Partenaires.objects.all()
    context = {
        'liste' : liste,
        'tenant' : request.tenant,
    }
    return render(request, 'tenant_folder/formations/liste_des_partenaires.html', context)


@login_required(login_url="institut_app:login")
def UpdatePartenaire(request, pk):
    partenaire = Partenaires.objects.get(id=pk)
    form = NewPartenaireForm(instance=partenaire)

    if request.method == "POST":
        form = NewPartenaireForm(request.POST, request.FILES ,instance=partenaire)
        if form.is_valid():
            updated_partenaire = form.save()
            UserActionLog.objects.create(
                user=request.user,
                action_type='UPDATE',
                target_model='Partenaire',
                target_id=str(updated_partenaire.id),
                details=f"Modification du partenaire : {updated_partenaire.nom} (Code: {updated_partenaire.code})",
                ip_address=request.META.get('REMOTE_ADDR')
            )
            if updated_partenaire.type_partenaire == 'etranger':
                instituts = Institut.objects.filter(is_visible=True).exclude(Q(schema_name='public') | Q(tenant_type='master'))
                for tenant in instituts:
                    with schema_context(tenant.schema_name):
                        try:
                            partenaire = Partenaires.objects.get(code=updated_partenaire.code)
                            partenaire.nom = updated_partenaire.nom
                            partenaire.adresse = updated_partenaire.adresse
                            partenaire.telephone = updated_partenaire.telephone
                            partenaire.email = updated_partenaire.email
                            partenaire.site_web = updated_partenaire.site_web
                            partenaire.type_partenaire = updated_partenaire.type_partenaire
                            partenaire.save()
                        except Partenaires.DoesNotExist:
                            pass
                messages.success(request, 'Partenaire mis à jour avec succès')
                return redirect('t_formations:listPartenaires')
            else:
                messages.success(request, 'Partenaire mis à jour avec succès')
                return redirect('t_formations:listPartenaires')
        else:
            messages.error(request, 'Une erreur s\'est produite lors du traitement de la requête')
            return redirect('t_formations:UpdatePartenaire', pk)

    context = {
        'form' : form,
        'tenant' : request.tenant,
    }
    return render(request,'tenant_folder/formations/update_partenaire.html',context)

@transaction.atomic
def addSpecialite(request):
    form = NewSpecialiteForm()
    if request.method == 'POST':
        form = NewSpecialiteForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Spécialité ajoutée avec succès")
            return redirect('t_formations:listSpecialites')
           
        else:
            errors = " | ".join(["{}: {}".format(field, ", ".join(errors)) for field, errors in form.errors.items()])
            messages.error(request, f"{errors}")
            return redirect('t_formations:addSpecialite')
        
    else:
        context = {
            'form' : form,
            'tenant' : request.tenant,
        }

        return render(request, 'tenant_folder/formations/nouvelle_specialite.html', context)

@login_required(login_url="institut_app:login")
@transaction.atomic
def updateSpecialite(request,pk):
    specialite = Specialites.objects.get(id = pk)
    form = NewSpecialiteForm(instance=specialite)
    if request.method == "POST":
        form = NewSpecialiteForm(request.POST, instance=specialite)
        if form.is_valid():
            updated_spec = form.save()
            updated_spec.save()

            formation  = specialite.formation
            formation.updated = True
            formation.save()

            messages.success(request, "Les données de la spécialité ont été mises à jour avec succès")
            return redirect('t_formations:detailSpecialite', pk)
        else:
            messages.error(request, "Une erreur s'est produite lors du traitement de la requête")
            return redirect('t_formations:detailSpecialite', pk)
    
    context = {
        'form' : form,
    }
    return render(request, "tenant_folder/formations/update_specialite.html", context)

def updateFraisInscription(request):
    pass

@login_required(login_url="institut_app:login")
def deleteFormation(request, pk):
    try:
        formation = Formation.objects.get(id=pk)
        
        if request.tenant.tenant_type != 'master' and formation.type_formation == 'etrangere':
            return JsonResponse({'success': False, 'message': "Vous ne pouvez pas supprimer une formation étrangère."})
            
        formation.delete()
        return JsonResponse({'success': True, 'message': "La formation a été supprimée avec succès."})
    except Formation.DoesNotExist:
        return JsonResponse({'success': False, 'message': "Formation introuvable."})
    except ProtectedError:
        return JsonResponse({'success': False, 'message': "Cette formation ne peut pas être supprimée car elle est liée à d'autres données (ex: inscriptions, promos)."})
    except Exception as e:
        return JsonResponse({'success': False, 'message': f"Une erreur est survenue: {str(e)}"})

@login_required(login_url="institut_app:login")
def deleteSpecialite(request, pk):
    try:
        specialite = Specialites.objects.get(id=pk)

        if request.tenant.tenant_type != 'master' and specialite.formation.type_formation == 'etranger':
            return JsonResponse({'success': False, 'message': "Vous ne pouvez pas supprimer cette spécialité."})
        
        specialite.delete()
        return JsonResponse({'success': True, 'message': "Spécialité supprimée avec succès."})
    except Specialites.DoesNotExist:
        return JsonResponse({'success': False, 'message': "Spécialité introuvable."})
    except ProtectedError:
        return JsonResponse({'success': False, 'message': "Cette spécialité ne peut pas être supprimée car elle est liée à d'autres données (ex: promotions, étudiants)."})
    except Exception as e:
        return JsonResponse({'success': False, 'message': f"Une erreur est survenue: {str(e)}"})
    
def deleteFraisInscription(request):
    pass

@login_required(login_url="institut_app:login")
def detailFormation(request, pk):
    formation = Formation.objects.get(id = pk)
    queryset = Specialites.objects.filter(formation = formation)
    if request.tenant.tenant_type != 'master':
        queryset = queryset.filter(is_visible=True)

    master_tenant = Institut.objects.filter(tenant_type='master').exclude(schema_name='public').first()
    
    entreprises = None
    if request.tenant.tenant_type == 'second':
        from institut_app.models import Entreprise
        entreprises = Entreprise.objects.all()

    context = {
        'formation' : formation,
        'tenant' : request.tenant,
        'specialite' : queryset,
        'master_tenant_name': master_tenant.nom if master_tenant else "l'établissement maître",
        'entreprises': entreprises,
    }
    return render(request, 'tenant_folder/formations/details_formation.html', context)

@login_required(login_url="institut_app:login")
def ApiUpdateFormationLegalEntity(request):
    if request.method != "POST":
        return JsonResponse({'success': False, 'message': 'Méthode non autorisée'})
    
    formation_id = request.POST.get('formation_id')
    entite_id = request.POST.get('entite_id')
    
    if not formation_id:
        return JsonResponse({'success': False, 'message': 'ID de formation manquant'})
        
    try:
        formation = Formation.objects.get(id=formation_id)
        
        # Security check: only allow if foreign formation and secondary tenant
        if formation.type_formation != 'etrangere' or request.tenant.tenant_type != 'second':
            return JsonResponse({'success': False, 'message': 'Action non autorisée pour ce type de formation ou d\'établissement.'})
            
        if entite_id:
            from institut_app.models import Entreprise
            try:
                entite = Entreprise.objects.get(id=entite_id)
                formation.entite_legal = entite
            except Entreprise.DoesNotExist:
                return JsonResponse({'success': False, 'message': 'Entité légale introuvable'})
        else:
            formation.entite_legal = None
            
        formation.save()
        return JsonResponse({'success': True, 'message': 'Entité légale mise à jour avec succès'})
    except Formation.DoesNotExist:
        return JsonResponse({'success': False, 'message': 'Formation introuvable'})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@login_required(login_url="institut_app:login")
@transaction.atomic
def ApiToggleTenantVisibility(request):
    if request.tenant.tenant_type != 'master':
        return JsonResponse({'success': False, 'message': 'Accès refusé. Seul le compte maître peut gérer la visibilité.'})
    
    tenant_id = request.POST.get('tenant_id')
    spec_code = request.POST.get('spec_code')
    
    if not tenant_id or not spec_code:
        return JsonResponse({'success': False, 'message': 'Paramètres manquants.'})

    try:
        inst = Institut.objects.get(id=tenant_id)
        with schema_context(inst.schema_name):
            try:
                spec_local = Specialites.objects.get(code=spec_code)
                spec_local.is_visible = not spec_local.is_visible
                spec_local.save()
                new_status = spec_local.is_visible
                return JsonResponse({
                    'success': True, 
                    'is_visible': new_status,
                    'message': f"Visibilité mise à jour pour {inst.nom}"
                })
            except Specialites.DoesNotExist:
                return JsonResponse({'success': False, 'message': 'Cette spécialité n\'est pas encore synchronisée sur cet institut.'})
                
    except Institut.DoesNotExist:
        return JsonResponse({'success': False, 'message': 'Institut introuvable.'})
    except Exception as e:
        return JsonResponse({'success': False, 'message': str(e)})

@login_required(login_url="institut_app:login")
def detailSpecialite(request, pk):
    object = Specialites.objects.get(id = pk)
    master_tenant = Institut.objects.filter(tenant_type='master').exclude(schema_name='public').first()

    context = {
        'object' : object,
        'tenant' : request.tenant,
        'master_tenant_name': master_tenant.nom if master_tenant else "l'établissement maître",
        'is_foreign_program': True if object.formation.type_formation == 'etrangere' or (object.formation.partenaire and object.formation.partenaire.type_partenaire == 'etranger') else False
    }

    if request.tenant.tenant_type == 'master':
        instituts = Institut.objects.filter(is_visible=True, tenant_type='second').order_by('nom')
        tenant_visibility = []
        master_module_count = Modules.objects.filter(specialite=object).count()
        for inst in instituts:
            try:
                with schema_context(inst.schema_name):
                    try:
                        spec_local = Specialites.objects.get(code=object.code)
                        is_visible = spec_local.is_visible
                        is_synced = True
                        module_count = Modules.objects.filter(specialite=spec_local).count()
                    except Specialites.DoesNotExist:
                        is_visible = False
                        is_synced = False
                        module_count = 0
                
                tenant_visibility.append({
                    'id': inst.id,
                    'nom': inst.nom,
                    'schema_name': inst.schema_name,
                    'is_visible': is_visible,
                    'is_synced': is_synced,
                    'module_count': module_count,
                    'master_module_count': master_module_count
                })
            except:
                pass
        
        context['tenant_visibility'] = tenant_visibility

    return render(request, 'tenant_folder/formations/details_specialite.html', context)

def ApiGetSpecialiteModule(request):
    try:
        id = request.GET.get('id')
        search_query = request.GET.get('search', '')
        page_number = int(request.GET.get('page', 1))
        page_size = int(request.GET.get('page_size', 10))

        if not id:
            return JsonResponse({'status': 'error', 'message': 'ID de spécialité manquant'}, status=400)

        # Base query
        queryset = Modules.objects.filter(specialite_id=id, is_archived=False)
        
        # Apply search filter
        if search_query:
            queryset = queryset.filter(
                Q(label__icontains=search_query) | 
                Q(code__icontains=search_query) |
                Q(code_interne__icontains=search_query)
            )
            
        # Order and values
        queryset = queryset.order_by('created_at')
        
        # All modules for select dropdown (unpaginated)
        all_modules = queryset.values('id', 'label')

        # Pagination
        paginator = Paginator(queryset.values('id', 'label','code','code_interne','coef','duree', 'est_valider'), page_size)
        page_obj = paginator.get_page(page_number)
        
        specialite = Specialites.objects.get(id=id)

        data = {
           'modules': list(page_obj.object_list),
           'all_modules': list(all_modules), # For the repartition select
           'nb_semestre': specialite.nb_semestre,
           'pagination': {
                'total_pages': paginator.num_pages,
                'current_page': page_obj.number,
                'total_items': paginator.count,
                'has_next': page_obj.has_next(),
                'has_previous': page_obj.has_previous(),
                'start_index': page_obj.start_index() if paginator.count > 0 else 0,
                'end_index': page_obj.end_index() if paginator.count > 0 else 0
            }
        }

        return JsonResponse(data, safe=False)
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

@login_required(login_url="institut_app:login")
@transaction.atomic
def ApiDuplicateSpecialite(request):
    if request.method != "POST":
        return JsonResponse({'success': False, 'message': 'Méthode non autorisée'})

    id_specialite = request.POST.get('id_specialite')
    version = request.POST.get('version')

    if not id_specialite:
        return JsonResponse({'success': False, 'message': 'ID de spécialité manquant'})

    try:
        source_spec = Specialites.objects.get(id=id_specialite)
    except Specialites.DoesNotExist:
        return JsonResponse({'success': False, 'message': 'Spécialité source introuvable'})

    # Generate new code
    new_version = version if version else "Copy"
    new_code = f"{source_spec.code}-{new_version}"
    
    # Check if code already exists
    if Specialites.objects.filter(code=new_code).exists():
        # Add a suffix if it exists
        count = 1
        while Specialites.objects.filter(code=f"{new_code}-{count}").exists():
            count += 1
        new_code = f"{new_code}-{count}"

    # Create new specialty
    new_spec = Specialites.objects.create(
        code=new_code,
        label=source_spec.label,
        prix=source_spec.prix,
        prix_double_diplomation=source_spec.prix_double_diplomation,
        duree=source_spec.duree,
        nb_semestre=source_spec.nb_semestre,
        branche=source_spec.branche,
        abr=source_spec.abr,
        formation=source_spec.formation,
        nb_tranche=source_spec.nb_tranche,
        responsable=source_spec.responsable,
        version=version,
        condition_access=source_spec.condition_access,
        updated_by=request.user
    )

    # Copy Modules
    source_modules = Modules.objects.filter(specialite=source_spec, is_archived=False)
    module_mapping = {} # To map old module ID to new module object for repartition copy

    for module in source_modules:
        old_id = module.id
        # Create new module
        new_module = Modules.objects.create(
            specialite=new_spec,
            code_interne=module.code_interne,
            label=module.label,
            duree=module.duree,
            coef=module.coef,
            n_elimate=module.n_elimate,
            systeme_eval=module.systeme_eval,
            created_by=request.user
        )
        module_mapping[old_id] = new_module

    # Copy ProgrammeFormation (Repartition)
    source_repartitions = ProgrammeFormation.objects.filter(specialite=source_spec)
    for repart in source_repartitions:
        if repart.module_id in module_mapping:
            ProgrammeFormation.objects.create(
                module=module_mapping[repart.module_id],
                specialite=new_spec,
                semestre=repart.semestre,
                created_by=request.user
            )

    return JsonResponse({
        'success': True, 
        'message': 'Spécialité dupliquée avec succès',
        'new_id': new_spec.id
    })

@login_required(login_url="institut_app:login")
def ApiGetRepartitionModule(request):
    id_specialite = request.GET.get('id_specialite')
    object = ProgrammeFormation.objects.filter(specialite = id_specialite).values('id', 'module__label','module__code','module__code_interne','semestre')

    return JsonResponse(list(object), safe=False)

@login_required(login_url="institut_app:login")
@transaction.atomic
def ApiDeleteCoursRepartition(request):
    try:
        id = request.GET.get('id')
        if not id:
            return JsonResponse({"success": False, "message": "Informations manquantes"})
        
        ProgrammeFormation.objects.get(id=id).delete()
        return JsonResponse({"success": True, "message": "L'affectation a été supprimée avec succès."})
    except ProgrammeFormation.DoesNotExist:
        return JsonResponse({"success": False, "message": "Affectation introuvable."})
    except ProtectedError:
        return JsonResponse({"success": False, "message": "Cette affectation ne peut pas être supprimée car elle est protégée."})
    except Exception as e:
        return JsonResponse({"success": False, "message": f"Erreur: {str(e)}"})

@login_required(login_url="institut_app:login")
def ApiAffectModuleSemestre(request):
    id_module = request.POST.get('id_module')
    semestre = request.POST.get('semestre')
    id_specialite = request.POST.get('id_specialite')
    try:
        specialite = Specialites.objects.get(id = id_specialite)
        module = Modules.objects.get(id = id_module)

        repartition = ProgrammeFormation(
           
            module = module,
            specialite  = specialite,
            semestre = semestre,
        )
        repartition.save()
        return JsonResponse({'success' : True, 'message' : "Le module a été affecté avec succès"})
    except Exception as e:
        return JsonResponse({'success' : False, 'message' : str(e)})

def ApiAddModule(request):

    label = request.POST.get('label')
    coef = request.POST.get('coef')
    duree = request.POST.get('duree')
    id = request.POST.get('id')
    code = request.POST.get('code_module')

    specialite = Specialites.objects.get(id = id)
    try:
        new_module = Modules.objects.create(
            label = label,
            coef = coef if coef != '' else None,
            duree = duree,
            specialite = specialite,
            code_interne = code,
            created_by = request.user,
        )

        new_module.save()
        return JsonResponse({'success' : True})
    except IntegrityError:
        return JsonResponse({'success' : False, 'message' : "Le module existe déjà"})

@login_required(login_url='institut_app:login')
def deleteModule(request):
    try:
        id = request.GET.get('id')
        obj = Modules.objects.get(id=id)
        obj.is_archived = True
        obj.save()
        return JsonResponse({'success': True, 'message': "Le module a été archivé avec succès."})
    except Modules.DoesNotExist:
        return JsonResponse({'success': False, 'message': "Module introuvable."})
    except Exception as e:
        return JsonResponse({'success': False, 'message': f"Erreur: {str(e)}"})

@login_required(login_url="institut_app:login")
def archiveModule(request):
    liste = Modules.objects.filter(is_archived = True)
    context = {
        'liste' : liste,
        'tenant' : request.tenant
    }
    return render(request, 'tenant_folder/formations/archive/archive_module.html', context)

@login_required(login_url="institut_app:login")
def ApiGetModuleDetails(request):
    id = request.GET.get('id')
    obj = Modules.objects.filter(id = id).values('id', 'code', 'label', 'duree', 'coef', 'n_elimate', 'systeme_eval', 'code_interne')
    return JsonResponse(list(obj), safe=False)

@login_required(login_url="institut_app:login")
@transaction.atomic
def ApiUpdateModule(request):

    id = request.POST.get('id')
    coef = request.POST.get('coef')
    duree = request.POST.get('duree')
    label = request.POST.get('label')
    code = request.POST.get('code')
    code_interne = request.POST.get('code_interne')
    n_elimate = request.POST.get('n_elimate')
    systeme_eval = request.POST.get('systeme_eval')

    module = Modules.objects.get(id= id)
    if code:
        module.code = code
    if code_interne:
        module.code_interne = code_interne
    else:
        # Fallback for old frontend behavior if needed
        module.code_interne = code
    module.duree = duree
    module.label = label
    module.coef = coef if coef != '' else None
    module.n_elimate = n_elimate if n_elimate != '' else None
    module.systeme_eval = systeme_eval

    # If the module was previously validated and we're making changes, we might want to unset the validation
    # This depends on business requirements, but typically modifications would require re-validation
    # Uncomment the next line if modifications should reset validation status
    # module.est_valider = False

    module.save()

    return JsonResponse({'success' : True, 'message' : "Les informations du module ont été mises à jour avec succès"})

def archiveFormation(request):
    pass
 
def detailModule(request):
    pass

def detailFraisInscription(request):
    pass

def listPromos(request):
    
    context = {
        
        'tenant' : request.tenant,
    }
    return render(request,'tenant_folder/formations/promos/list_promos.html',context)

@login_required(login_url="institut_app:login")
def ApiListePromos(request):
    liste= Promos.objects.filter().values('id','label','session','etat','code','begin_year','end_year','annee_academique')

    for l in liste:
        l_obj = Promos.objects.get(id = l['id'])
        l['etat_label'] = l_obj.get_etat_display()
        l['session_label'] = l_obj.get_session_display()

    return JsonResponse(list(liste), safe=False)

@login_required(login_url="institut_app:login")
def ApiCheckPromoCode(request):
    code = request.GET.get('code')
    exclude_id = request.GET.get('exclude_id')
    
    if not code:
        return JsonResponse({'exists': False})
        
    query = Promos.objects.filter(code__iexact=code)
    if exclude_id:
        query = query.exclude(id=exclude_id)
        
    exists = query.exists()
    return JsonResponse({'exists': exists})

def ApiListeFormation(request):
    liste = Formation.objects.all().values('id', 'nom')
    return JsonResponse(list(liste), safe=False)

def ApiListeSpecialiteByFormation(request):
    id= request.GET.get('id_formation')
    formation_obj = Formation.objects.get(id = id)
    liste = Specialites.objects.filter(formation = formation_obj.code).values('id', 'label','code')
    return JsonResponse(list(liste), safe=False)

@login_required(login_url="institut_app:login")
@transaction.atomic
def AddPromo(request):
    form = PromoForm()
    if request.method =="POST":
        post_data = request.POST.copy()
        if 'label' in post_data: post_data['label'] = post_data['label'].upper()
        if 'code' in post_data: post_data['code'] = post_data['code'].upper()
        form = PromoForm(post_data)
        if form.is_valid():
            form.save()
            messages.success(request, 'Promo ajoutée avec succès')
            return redirect('t_formations:listPromos')
        else:
            messages.error(request, 'Une erreur s\'est produite lors du traitement de la requête')
            return redirect('t_formations:AddPromo')
    context = {
        'form' : form,
        'tenant' : request.tenant
    }
    return render(request, 'tenant_folder/formations/promos/new_promo.html', context)

@login_required(login_url="institut_app:login")
def detailPromo(request, pk):
    promo = Promos.objects.get(id=pk)
    groupes = Groupe.objects.filter(promotion=promo).select_related('specialite')
    
    from t_crm.models import FicheDeVoeux, FicheVoeuxDouble
    
    # 1. Standard prospects by specialty
    fiches_standard = FicheDeVoeux.objects.filter(promo=promo).select_related('prospect', 'specialite')
    
    # Group by specialty
    specialites_dict = {}
    for fiche in fiches_standard:
        if fiche.specialite:
            spec_id = fiche.specialite.id
            if spec_id not in specialites_dict:
                specialites_dict[spec_id] = {
                    'specialite': fiche.specialite,
                    'fiches': []
                }
            specialites_dict[spec_id]['fiches'].append(fiche)
            
    # Convert dict to sorted list of specialties
    specialites_standard = sorted(specialites_dict.values(), key=lambda x: x['specialite'].label or '')

    # 2. Double diplomation prospects by double diplomation combination
    fiches_double = FicheVoeuxDouble.objects.filter(promo=promo).select_related('prospect', 'specialite')
    
    # Group by double diplomation Combination
    double_dict = {}
    for fiche in fiches_double:
        if fiche.specialite: # This is a DoubleDiplomation object
            double_id = fiche.specialite.id
            if double_id not in double_dict:
                double_dict[double_id] = {
                    'double_diplomation': fiche.specialite,
                    'fiches': []
                }
            double_dict[double_id]['fiches'].append(fiche)
            
    specialites_double = sorted(double_dict.values(), key=lambda x: x['double_diplomation'].label or '')

    total_prospects_standard = len(fiches_standard)
    total_prospects_double = len(fiches_double)
    
    context = {
        'promo': promo,
        'groupes': groupes,
        'tenant': request.tenant,
        'specialites_standard': specialites_standard,
        'specialites_double': specialites_double,
        'total_prospects_standard': total_prospects_standard,
        'total_prospects_double': total_prospects_double,
        'total_prospects_all': total_prospects_standard + total_prospects_double
    }
    return render(request, 'tenant_folder/formations/promos/details_promo.html', context)

@login_required(login_url="institut_app:login")
def ApiDeletePromo(request):
    id = request.POST.get('id')
    promo = Promos.objects.get(id = id)
    if promo.etat == 'active':
        return JsonResponse({'success' : False, 'message' : 'Vous ne pouvez pas supprimer une promo active'})
    else:
        promo.delete()
        return JsonResponse({'success' : True, 'message' : 'Promo supprimée avec succès'})

@login_required(login_url="institut_app:login")
def ApiGetPromo(request):
    id = request.GET.get('id')
    promo = Promos.objects.filter(id = id).values('id', 'label', 'session','code','begin_year','end_year','annee_academique')
    return JsonResponse(list(promo), safe=False)

@login_required(login_url="institut_app:login")
@transaction.atomic
def ApiUpdatePromo(request):
    if request.method == "POST":
        id = request.POST.get('id')
        label = request.POST.get('label', '').upper()
        session = request.POST.get('session')
        new_code  = request.POST.get('new_code', '').upper()
        new_begin_year = request.POST.get('new_begin_year')
        new_end_year = request.POST.get('new_end_year')

        annee_academique = request.POST.get('annee_academique')

        if not label or not session:
            return JsonResponse({'success' : False, 'message' : 'Veuillez remplir les champs obligatoires'})
        else:
            promo = Promos.objects.get(id = id)
            promo.label = label
            promo.session = session
            promo.code = new_code
            promo.begin_year = new_begin_year
            promo.end_year = new_end_year
            promo.annee_academique = annee_academique
            promo.save()
            return JsonResponse({'success' : True, 'message' : 'Promo mise à jour avec succès'})
    else:
        return JsonResponse({"status" : "error"})
    
def ApiActivatePromo(request):
    id = request.POST.get('id')
    promo = Promos.objects.get(id = id)
    promo.etat = 'active'
    promo.save()
    return JsonResponse({'success' : True, 'message' : 'Promo activée avec succès'})

def SpecialitePromo(request):
    
    context = {
        'tenant' : request.tenant
    }
    return render(request, 'tenant_folder/formations/promos/specialite_promo.html', context)

def PageListeModules(request):
    return render(request, 'tenant_folder/formations/modules/modules.html', {'tenant' : request.tenant})

def ApiGetModules(request):
    try:
        # Get parameters
        page_number = int(request.GET.get('page', 1))
        page_size = int(request.GET.get('page_size', 10))
        search_query = request.GET.get('search', '')
        status_filter = request.GET.get('status', '0')
        order_by = request.GET.get('order_by', '-created_at')
        
        # Base query
        queryset = Modules.objects.filter(is_archived=False)
        
        # Apply filters
        if search_query:
            queryset = queryset.filter(
                Q(label__icontains=search_query) | 
                Q(code__icontains=search_query) |
                Q(code_interne__icontains=search_query)
            )
        
        if status_filter == 'actif':
            queryset = queryset.filter(est_valider=True)
        elif status_filter == 'inactif':
            queryset = queryset.filter(est_valider=False)
            
        # Apply ordering
        if order_by in ['created_at', '-created_at']:
            queryset = queryset.order_by(order_by)
        else:
            queryset = queryset.order_by('-created_at')

        # Global stats (for KPIs)
        total_count = Modules.objects.filter(is_archived=False).count()
        valides_count = Modules.objects.filter(is_archived=False, est_valider=True).count()
        en_attente_count = Modules.objects.filter(is_archived=False, est_valider=False).count()
            
        # Pagination
        paginator = Paginator(queryset.values('id', 'label', 'coef', 'duree', 'code', 'est_valider', 'created_at', 'specialite__label'), page_size)
        page_obj = paginator.get_page(page_number)
        
        data = {
            'data': list(page_obj.object_list),
            'pagination': {
                'total_pages': paginator.num_pages,
                'current_page': page_obj.number,
                'total_items': paginator.count,
                'has_next': page_obj.has_next(),
                'has_previous': page_obj.has_previous(),
                'start_index': page_obj.start_index() if paginator.count > 0 else 0,
                'end_index': page_obj.end_index() if paginator.count > 0 else 0
            },
            'stats': {
                'total': total_count,
                'valides': valides_count,
                'en_attente': en_attente_count,
                'inactifs': 0 # Assuming 'archived' means inactive or we define it differently
            }
        }
        
        return JsonResponse(data, safe=False)
        
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

def ApiGetModuleDetails(request):
    id = request.GET.get('id')
    
    if not id:
        return JsonResponse({'status': 'error', 'message': 'ID du module est requis'}, safe=False)
    
    try:
        details = Modules.objects.get(id=id)
        # Get associated trainers for this module
        associated_trainers = EnseignantModule.objects.filter(module=details).select_related('formateur')
        trainers_list = []
        
        for association in associated_trainers:
            trainer = association.formateur
            trainers_list.append({
                'id': trainer.id,
                'nom': trainer.nom,
                'prenom': trainer.prenom,
                'email': trainer.email,
                'telephone': trainer.telephone,
                'diplome': trainer.diplome,
            })
        
        data = {
            'id': details.id,
            'code': details.code,
            'code_interne': details.code_interne,
            'label': details.label,
            'duree': details.duree,
            'coef': details.coef,
            'n_elimate': details.n_elimate,
            'systeme_eval': details.systeme_eval,
            'specialite_label': details.specialite.label if details.specialite else 'N/A',
            'associated_trainers': trainers_list
        }
        
        return JsonResponse(data, safe=False)
    except Modules.DoesNotExist:
        return JsonResponse({'status': 'error', 'message': 'Module non trouvé'}, safe=False)


def get_module_details_with_teachers(request):
    module_id = request.GET.get('module_id')
    if module_id:
        try:
            # Get module details
            module = Modules.objects.get(id=module_id)
            
            # Get associated trainers for this module
            teacher_modules = EnseignantModule.objects.filter(module_id=module_id).select_related('formateur')
            teachers_data = []
            
            for tm in teacher_modules:
                formateur = tm.formateur
                teachers_data.append({
                    'id': formateur.id,
                    'nom': formateur.nom,
                    'prenom': formateur.prenom,
                    'telephone': formateur.telephone,
                    'email': formateur.email,
                    'diplome': formateur.diplome,
                })
            
            # Prepare module data
            module_data = {
                'id': module.id,
                'code': module.code or '',
                'code_interne': module.code_interne or '',
                'label': module.label or '',
                'duree': module.duree or '',
                'coef': module.coef or '',
                'n_elimate': module.n_elimate or '',
                'systeme_eval': module.systeme_eval or '',
                'teachers': teachers_data
            }
            
            return JsonResponse(module_data, safe=False)
        except Modules.DoesNotExist:
            return JsonResponse({'status': 'error', 'message': 'Module non trouvé'})
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)})
    else:
        return JsonResponse({'status': 'error', 'message': 'ID du module requis'})

@login_required(login_url='intitut_app:login')
def ApiLoadDocuments(request):
    code_formation = request.GET.get('code_formation')
    documents = DossierInscription.objects.filter(formation__code = code_formation).values('id','label', 'is_required', 'include_in_tracking')


    return JsonResponse(list(documents), safe=False)

@login_required(login_url="institut_app:login")
def ApiAddDocument(request):
    label = request.POST.get('nom_doc')
    required = request.POST.get('required')
    code_formation = request.POST.get('code_formation')
    
    if required == "false":
        _required = False
    else:
        _required = True

    include_in_tracking = request.POST.get('include_in_tracking')
    
    if include_in_tracking == "false":
        _include = False
    else:
        _include = True

    DossierInscription.objects.create(
        formation = Formation.objects.get(code = code_formation),
        label = label,
        is_required = _required,
        include_in_tracking = _include
    )

    return JsonResponse({"status" : "success","message" : "Le document a été ajouté avec succès"})

@login_required(login_url="institu_app:login")
@transaction.atomic
def ApiDeleteDoc(request):
    id_doc = request.POST.get('id_doc')
    if id_doc:
        obj = DossierInscription.objects.get(id = id_doc)
        obj.delete()
        return JsonResponse({'status' : 'success', 'message' : 'Document supprimé avec succès' })
    else:
        return JsonResponse({'status' : 'error', 'message' : 'Une erreur est survenue lors du traitement de la requête' })


@login_required(login_url="institut_app:login")
@transaction.atomic
def ApiUpdateDoc(request):
    id_doc = request.POST.get('id_doc')
    label = request.POST.get('nom_doc')
    required = request.POST.get('required')

    if not id_doc or not label:
        return JsonResponse({'status' : 'error', 'message' : 'Informations manquantes' })

    try:
        doc = DossierInscription.objects.get(id=id_doc)
        doc.label = label
        doc.is_required = (required == "true")
        
        include_in_tracking = request.POST.get('include_in_tracking')
        if include_in_tracking is not None:
             doc.include_in_tracking = (include_in_tracking == "true")
             
        doc.save()
        return JsonResponse({'status' : 'success', 'message' : 'Document mis à jour avec succès' })
    except DossierInscription.DoesNotExist:
        return JsonResponse({'status' : 'error', 'message' : 'Document introuvable' })
    except Exception as e:
        return JsonResponse({'status' : 'error', 'message' : str(e) })
    

@login_required(login_url="institut_app:login")
def ApiLoadSpecForPartenaire(request):
    if request.method == "GET":
        code_partenaire = request.GET.get('id_partenaire')
        print(code_partenaire)
        liste = Specialites.objects.filter(formation__partenaire__code = code_partenaire).values('id', 'label','code','version')
        return JsonResponse(list(liste),safe=False)


# Vue pour enregistrer les documents d'impression sélectionnés
from django.views.decorators.csrf import csrf_exempt
import json

@login_required(login_url="institut_app")
def ApiLoadDocumentTemplate(request):
    formation_id = request.GET.get('formation_id')

    formation = Formation.objects.get(id=formation_id)

    all_docs = DocumentTemplate.objects.filter(is_active = True)
    selected_docs = formation.documents.values_list('id', flat=True)

    data = {
        "documents": [
            {
                "id": doc.id,
                "title": doc.title,
                "checked": doc.id in selected_docs
            }
            for doc in all_docs
        ]
    }

    return JsonResponse(data)

@csrf_exempt
@login_required
def save_print_documents(request, pk):
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            formation_id = data.get('formation_id')
            documents = data.get('documents', [])  # liste d'IDs

            # Récupérer la formation
            formation = Formation.objects.get(id=formation_id)

            # Associer les documents (remplace les anciens)
            formation.documents.set(documents)

            return JsonResponse({
                'success': True,
                'message': f'{len(documents)} documents enregistrés pour la formation {formation.nom}'
            })
        except Exception as e:
            return JsonResponse({
                'success': False,
                'error': str(e)
            }, status=500)
    else:
        return JsonResponse({
            'success': False,
            'error': 'Méthode non autorisée'
        }, status=405)

# Vue pour récupérer les documents d'impression sélectionnés
@login_required
def get_print_documents(request, pk):
    if request.method == 'GET':
        try:
            formation = Formation.objects.get(id=pk)

            selected_documents = list(formation.documents.values_list('id', flat=False))
            
            return JsonResponse({
                'success': True,
                'documents': selected_documents
            })
        except Exception as e:
            return JsonResponse({
                'success': False,
                'error': str(e)
            }, status=500)
    else:
        return JsonResponse({
            'success': False,
            'error': 'Méthode non autorisée'
        }, status=405)

    

@login_required(login_url="institut_app:login")
def import_data_view(request):
    if request.method == 'POST':
        # Check if it's the confirmation step
        if 'confirm_import' in request.POST:
            data_json = request.session.get('import_data')
            data_type = request.session.get('import_data_type')
            
            if not data_json or not data_type:
                messages.error(request, "Session expirée ou données manquantes. Veuillez recommencer.")
                return redirect('t_formations:import_data_view')
            
            try:
                # data_json is a JSON string of list of dicts or list of dicts directly logic?
                # session stores python types usually, so list of dicts.
                count = import_data(data_json, data_type, user=request.user)
                messages.success(request, f"{count} enregistrements importés/mis à jour avec succès.")
                
                # Clear session
                del request.session['import_data']
                del request.session['import_data_type']
                
                return redirect('t_formations:import_data_view')
            except Exception as e:
                messages.error(request, f"Erreur lors de l'import : {str(e)}")
                return redirect('t_formations:import_data_view')
        
        # Upload Step
        form = ImportDataForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                file = request.FILES['file']
                data_type = form.cleaned_data['data_type']
                
                parsed_data = handle_uploaded_file(file)
                report = verify_data(parsed_data, data_type)
                
                # Store valid rows for confirmation (if any)
                # We store only the valid rows to be imported.
                # If there are errors, user might want to fix file and re-upload, or proceed with valid rows?
                # Let's assume we allow proceeding with valid rows.
                
                if report['valid_rows']:
                     request.session['import_data'] = report['valid_rows']
                     request.session['import_data_type'] = data_type
                
                context = {
                    'report': report,
                    'data_type': data_type,
                    'tenant': request.tenant,
                    'has_valid_rows': len(report['valid_rows']) > 0
                }
                return render(request, 'tenant_folder/formations/import_preview.html', context)
                
            except ValidationError as e:
                messages.error(request, e.message)
            except Exception as e:
                messages.error(request, f"Erreur inattendue : {str(e)}")
    else:
        form = ImportDataForm()

    return render(request, 'tenant_folder/formations/import_data.html', {'form': form, 'tenant': request.tenant})


@login_required(login_url="institut_app:login")
def export_formations(request):
    format_type = request.GET.get('format', 'csv')
    formations = Formation.objects.prefetch_related('formation_specilite').all()

    # Define headers
    headers = [
        'Code Formation', 'Nom Formation', 'Description Formation', 'Durée Formation', 'Partenaire', 'Type de formation', 'Qualification Formation', 'Frais Inscription', 'Prix Formation',
        'Code Spécialité', 'Label Spécialité', 'Durée Spécialité', 'Branche', 'Prix Spécialité', 'Version Spécialité', 'Semestres Spécialité', 'Tranches Spécialité', 'Prix Double Diplomation', 'Abréviation Spécialité', "Conditions d'accès",
        'Code Module', 'Code Interne Module', 'Label Module', 'Durée Module', 'Coefficient', 'Semestre Module'
    ]

    def get_rows_for_formation(f):
        rows = []
        partenaire_val = f.partenaire.code if f.partenaire else ''
        base_f_row = [
            f.code, f.nom, f.description, f.duree, partenaire_val,
            f.type_formation if f.type_formation else '',
            f.qualification if f.qualification else '',
            f.frais_inscription, f.prix_formation
        ]
        
        specialites = f.formation_specilite.all()
        if not specialites:
            rows.append(base_f_row + [''] * 12 + [''] * 6)
        else:
            for spec in specialites:
                base_s_row = [
                    spec.code, spec.label, spec.duree, spec.branche, spec.prix,
                    spec.version, spec.nb_semestre, spec.nb_tranche, spec.prix_double_diplomation, spec.abr, spec.condition_access
                ]
                
                modules = Modules.objects.filter(specialite=spec)
                if not modules:
                    rows.append(base_f_row + base_s_row + [''] * 6)
                else:
                    for mod in modules:
                        programme = ProgrammeFormation.objects.filter(module=mod, specialite=spec).first()
                        semestre = programme.semestre if programme else ''
                        mod_row = [
                            mod.code, mod.code_interne, mod.label, mod.duree, mod.coef, semestre
                        ]
                        rows.append(base_f_row + base_s_row + mod_row)
        return rows


    if format_type == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="formations_export.xlsx"'

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Formations'
        sheet.append(headers)

        for f in formations:
            for row in get_rows_for_formation(f):
                sheet.append(row)

        workbook.save(response)
        return response

    else:
        # Default to CSV
        response = HttpResponse(content_type='text/csv')
        response.write('\ufeff'.encode('utf8'))
        response['Content-Disposition'] = 'attachment; filename="formations_export.csv"'

        writer = csv.writer(response, delimiter=';')
        writer.writerow(headers)

        for f in formations:
            for row in get_rows_for_formation(f):
                writer.writerow(row)

        return response

@login_required(login_url="institut_app:login")
def export_partenaires(request):
    format_type = request.GET.get('format', 'csv')
    partenaires = Partenaires.objects.all()

    UserActionLog.objects.create(
        user=request.user,
        action_type='EXPORT',
        target_model='Partenaire',
        details=f"Exportation de la liste des partenaires (Format: {format_type})",
        ip_address=request.META.get('REMOTE_ADDR')
    )

    if format_type == 'excel':
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="partenaires_export.xlsx"'

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Partenaires'

        # Headers
        headers = ['Code', 'Nom', 'Adresse', 'Téléphone', 'Email', 'Site Web', 'Type de partenaire', 'Etat']
        sheet.append(headers)

        for p in partenaires:
            sheet.append([
                p.code,
                p.nom,
                p.adresse,
                p.telephone,
                p.email,
                p.site_web,
                p.type_partenaire if p.type_partenaire else '',
                p.etat
            ])

        workbook.save(response)
        return response

    else:
        # Default to CSV
        response = HttpResponse(content_type='text/csv')
        # Add UTF-8 BOM for Excel compatibility with CSV
        response.write('\ufeff'.encode('utf8'))
        response['Content-Disposition'] = 'attachment; filename="partenaires_export.csv"'

        writer = csv.writer(response, delimiter=';')
        writer.writerow(['Code', 'Nom', 'Adresse', 'Téléphone', 'Email', 'Site Web', 'Type de partenaire', 'Etat'])

        for p in partenaires:
            writer.writerow([
                p.code,
                p.nom,
                p.adresse,
                p.telephone,
                p.email,
                p.site_web,
                p.type_partenaire if p.type_partenaire else '',
                p.etat
            ])

        return response
